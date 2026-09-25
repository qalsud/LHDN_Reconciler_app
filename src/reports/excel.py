"""Multi-tab formatted Excel audit workbook exporter (OpenPyXL)."""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from loguru import logger
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from src.engine.reconciler import ReconcileResult

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=10)
TITLE_FONT = Font(name="Calibri", bold=True, color="1F4E78", size=14)
SUBTITLE_FONT = Font(name="Calibri", color="595959", size=9)
THIN_BORDER = Border(
    left=Side(style="thin", color="BFBFBF"),
    right=Side(style="thin", color="BFBFBF"),
    top=Side(style="thin", color="BFBFBF"),
    bottom=Side(style="thin", color="BFBFBF"),
)
MISSING_UUID_FILL = PatternFill("solid", fgColor="FFC7CE")  # light red
SST_MISMATCH_FILL = PatternFill("solid", fgColor="FFEB9C")  # light yellow
MATCHED_FILL = PatternFill("solid", fgColor="C6EFCE")  # light green
ZEBRA_FILL = PatternFill("solid", fgColor="F2F2F2")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)

SHEET_TITLES = {
    "Matched": "Matched",
    "Unsubmitted_Sales": "Unsubmitted_Sales",
    "SST_Rate_Mismatch": "SST_Rate_Mismatch",
    "Missing_UUID": "Missing_UUID",
}

COLUMN_HEADERS = [
    "GL Date", "LHDN Date", "TIN",
    "GL Reference", "LHDN Reference", "LHDN UUID",
    "Subtotal (RM)", "SST GL (RM)", "SST LHDN (RM)", "SST Variance (RM)",
    "Total GL (RM)", "Total LHDN (RM)", "Total Variance (RM)",
    "Date Diff (days)", "Match Stage",
    "Anomaly (0-100)", "AI Narrative",
]
COLUMN_KEYS = [
    "transaction_date", "invoice_date_lhdn", "tin",
    "invoice_reference_gl", "invoice_reference_lhdn", "lhdn_uuid",
    "subtotal_gl", "sst_gl", "sst_lhdn", "sst_variance",
    "total_gl", "total_lhdn", "total_variance",
    "date_diff_days", "match_stage",
    "anomaly_score", "ai_narrative",
]
MONEY_KEYS = {"subtotal_gl", "sst_gl", "sst_lhdn", "sst_variance", "total_gl", "total_lhdn", "total_variance"}


def _style_header(ws, ncols: int) -> None:
    for col in range(1, ncols + 1):
        cell = ws.cell(row=3, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER
        cell.border = THIN_BORDER


def _write_bucket_sheet(
    wb: Workbook,
    title: str,
    df: pd.DataFrame,
    row_fill: PatternFill | None = None,
    note: str = "",
) -> None:
    ws = wb.create_sheet(title=title)
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.sheet_view.showGridLines = False
    ws["A1"] = title.replace("_", " ")
    ws["A1"].font = TITLE_FONT
    ws["A2"] = note or f"{len(df)} record(s)"
    ws["A2"].font = SUBTITLE_FONT

    for col_idx, header in enumerate(COLUMN_HEADERS, start=1):
        ws.cell(row=3, column=col_idx, value=header)
    _style_header(ws, len(COLUMN_HEADERS))

    for r_idx, (_, row) in enumerate(df.iterrows(), start=4):
        for c_idx, key in enumerate(COLUMN_KEYS, start=1):
            value = row.get(key)
            if pd.isna(value):
                value = "" if key not in MONEY_KEYS else None
            if key in ("transaction_date", "invoice_date_lhdn") and value not in ("", None):
                try:
                    value = pd.Timestamp(value).date()
                except (ValueError, TypeError):
                    pass
            cell = ws.cell(row=r_idx, column=c_idx, value=value)
            cell.font = Font(name="Calibri", size=9)
            cell.border = THIN_BORDER
            cell.alignment = CENTER if c_idx not in (4, 5, 6, 17) else LEFT
            if key == "anomaly_score" and isinstance(value, (int, float)):
                cell.number_format = '0.0'
            if key in MONEY_KEYS and isinstance(value, (int, float)):
                cell.number_format = '#,##0.00'
            # Variance emphasis
            if key == "sst_variance" and isinstance(value, (int, float)) and abs(value) > 0.051:
                cell.fill = SST_MISMATCH_FILL
                cell.font = Font(name="Calibri", size=9, bold=True)
            elif row_fill is not None and (r_idx % 2 == 0 or title in ("Missing_UUID", "SST_Rate_Mismatch")):
                # Full-row tint for exception buckets; zebra for the rest.
                cell.fill = row_fill
            elif r_idx % 2 == 0:
                cell.fill = ZEBRA_FILL

    widths = [12, 12, 16, 16, 16, 38, 13, 12, 13, 14, 13, 14, 15, 12, 15, 12, 45]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A4"
    ws.auto_filter.ref = f"A3:{get_column_letter(len(COLUMN_HEADERS))}{max(3, len(df) + 3)}"


def _write_summary(wb: Workbook, result: ReconcileResult) -> None:
    ws = wb.active
    ws.title = "Summary"
    ws.sheet_view.showGridLines = False
    s = result.summary

    ws["A1"] = "LHDN MyInvois vs General Ledger — Reconciliation Summary"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = (
        f"Date tolerance ±{s['date_tolerance_days']} day(s)  |  "
        f"Amount tolerance ±RM {s['amount_tolerance_rm']:.2f}  |  "
        f"SST tolerance ±RM {s['sst_tolerance_rm']:.2f}"
    )
    ws["A2"].font = SUBTITLE_FONT

    headers = ["Bucket", "Records", "Share of GL", "Interpretation"]
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=4, column=c, value=h)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER
        cell.border = THIN_BORDER

    gl_total = max(int(s["gl_total"]), 1)
    rows = [
        ("GL records (input)", s["gl_total"], "—", "Ledger sales population"),
        ("LHDN documents (input)", s["lhdn_total"], "—", "MyInvois submission population"),
        ("Matched", s["matched"], f"{s['matched']/gl_total:.1%}", "Fully reconciled"),
        ("Unsubmitted_Sales", s["unsubmitted_sales"], f"{s['unsubmitted_sales']/gl_total:.1%}", "In GL, absent from MyInvois"),
        ("SST_Rate_Mismatch", s["sst_rate_mismatch"], f"{s['sst_rate_mismatch']/gl_total:.1%}", "Counterpart found, tax differs"),
        ("Missing_UUID", s["missing_uuid"], "—", "UUID unlinked / LHDN-only docs"),
    ]
    fills = [None, None, MATCHED_FILL, None, SST_MISMATCH_FILL, MISSING_UUID_FILL]
    for r, (bucket, count, share, interp) in enumerate(rows, start=5):
        for c, val in enumerate([bucket, count, share, interp], start=1):
            cell = ws.cell(row=r, column=c, value=val)
            cell.font = Font(name="Calibri", size=10, bold=(c == 1))
            cell.border = THIN_BORDER
            cell.alignment = CENTER if c in (2, 3) else LEFT
            if fills[r - 5] is not None:
                cell.fill = fills[r - 5]

    for col, width in zip("ABCD", [26, 12, 14, 38]):
        ws.column_dimensions[col].width = width


def export_workbook(result: ReconcileResult, output_path: str | Path) -> Path:
    """Write the multi-tab audit workbook. Returns the resolved output path."""
    out = Path(output_path)
    try:
        out.parent.mkdir(parents=True, exist_ok=True)
    except OSError as exc:
        raise OSError(f"Cannot create output directory '{out.parent}': {exc}") from exc

    buckets = result.buckets()
    try:
        wb = Workbook()
        _write_summary(wb, result)
        _write_bucket_sheet(wb, "Matched", buckets["Matched"], row_fill=None,
                            note=f"{len(buckets['Matched'])} fully reconciled record(s).")
        _write_bucket_sheet(wb, "Unsubmitted_Sales", buckets["Unsubmitted_Sales"], row_fill=None,
                            note=f"{len(buckets['Unsubmitted_Sales'])} GL sale(s) with no MyInvois counterpart — submit urgently.")
        _write_bucket_sheet(wb, "SST_Rate_Mismatch", buckets["SST_Rate_Mismatch"], row_fill=SST_MISMATCH_FILL,
                            note=f"{len(buckets['SST_Rate_Mismatch'])} record(s) where SST/total differs beyond tolerance (yellow).")
        _write_bucket_sheet(wb, "Missing_UUID", buckets["Missing_UUID"], row_fill=MISSING_UUID_FILL,
                            note=f"{len(buckets['Missing_UUID'])} record(s) with unlinked/missing LHDN UUID (red).")
        wb.save(out)
    except Exception as exc:
        raise OSError(f"Failed to write Excel workbook '{out}': {exc}") from exc

    logger.info("Wrote reconciliation workbook: {} ({} sheets)", out, len(wb.sheetnames))
    return out
