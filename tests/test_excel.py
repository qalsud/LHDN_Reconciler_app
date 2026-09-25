"""Tests for the Excel audit exporter."""

from pathlib import Path

from openpyxl import load_workbook

from src.engine.reconciler import ReconcileConfig, reconcile
from src.parsers.gl_parser import parse_gl_csv
from src.parsers.lhdn_parser import parse_lhdn_json
from src.reports.excel import export_workbook


def test_export_workbook_sheets_and_highlight(tmp_path):
    gl = parse_gl_csv("samples/gl_sample.csv")
    lhdn = parse_lhdn_json("samples/lhdn_sample.json")
    result = reconcile(gl, lhdn, ReconcileConfig())
    out = export_workbook(result, tmp_path / "summary.xlsx")

    assert Path(out).is_file()
    wb = load_workbook(out)
    assert wb.sheetnames == ["Summary", "Matched", "Unsubmitted_Sales", "SST_Rate_Mismatch", "Missing_UUID"]

    # Summary bucket counts mirror the engine output.
    summary = wb["Summary"]
    assert summary["B7"].value == 6   # Matched
    assert summary["B8"].value == 3   # Unsubmitted_Sales
    assert summary["B9"].value == 2   # SST_Rate_Mismatch
    assert summary["B10"].value == 2  # Missing_UUID

    # Exception buckets carry their tint on the first data row.
    mismatch_fill = wb["SST_Rate_Mismatch"]["A4"].fill.start_color.rgb
    assert mismatch_fill == "00FFEB9C"
    missing_fill = wb["Missing_UUID"]["A4"].fill.start_color.rgb
    assert missing_fill == "00FFC7CE"

    # Phase 4 columns present on every bucket sheet.
    headers = [wb[name]["A3"].value for name in wb.sheetnames[1:]]
    assert all(h == "GL Date" for h in headers)
    for name in wb.sheetnames[1:]:
        ws = wb[name]
        row_headers = [ws.cell(row=3, column=c).value for c in range(1, ws.max_column + 1)]
        assert "Anomaly (0-100)" in row_headers
        assert "AI Narrative" in row_headers
